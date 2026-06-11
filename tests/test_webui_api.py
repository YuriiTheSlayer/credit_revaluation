"""Тесты Python-API JS-фронтенда: загрузка, банк/веса/фильтры, таблица,
маппинги, экспорт — без открытия окна (вся логика живёт в webui/api.py)."""

import json

import openpyxl
import pytest

from webui.api import Api


@pytest.fixture
def api(tmp_path, kniga_path, sales_path) -> Api:
    a = Api(mapping_path=tmp_path / "mapping.json",
            override_path=tmp_path / "override.json")
    snap = a.load_competitors(str(kniga_path))
    assert "error" not in snap
    snap = a.load_sales(str(sales_path))
    assert "error" not in snap
    return a


def _row(snap, sku):
    rows = {r["sku"]: r for r in snap["table"]["rows"]}
    assert sku in rows, f"SKU {sku} нет на текущей странице"
    return rows[sku]


def test_boot_before_load(tmp_path):
    a = Api(mapping_path=tmp_path / "m.json", override_path=tmp_path / "o.json")
    snap = a.boot()
    assert snap["loaded"] is False
    assert snap["table"] is None and snap["kpi"] == []
    assert snap["version"]                     # версия видна в шапке UI
    json.dumps(snap, allow_nan=False)          # снапшот строго JSON-чистый


def test_snapshot_after_load(api):
    snap = api.boot()
    assert snap["loaded"] is True
    assert snap["banks"] == ["Monobank", "ПриватБанк", "ПУМБ"]
    assert snap["bank"] == "Monobank"
    assert snap["hasSales"] is True and snap["weight"] == "sales"
    assert len(snap["kpi"]) >= 4
    assert snap["table"]["total"] == 12
    labels = [c["label"] for c in snap["table"]["columns"]]
    assert "Предл. КМ" in labels and "Откл." in labels
    assert snap["charts"]["terms"]["series"][0]["name"] == "Comfy"
    assert "Конкуренты:" not in (snap["statuses"]["competitors"] or "")
    json.dumps(snap, allow_nan=False)          # NaN не должны утекать в JS


def test_bank_switch_changes_numbers(api):
    mono = _row(api.boot(), "20671")
    assert mono["pay::foxtrot.com.ua"] == 6
    privat = _row(api.set_bank("ПриватБанк"), "20671")
    assert privat["pay::foxtrot.com.ua"] == 15


def test_weight_requires_sales(tmp_path, kniga_path):
    a = Api(mapping_path=tmp_path / "m.json", override_path=tmp_path / "o.json")
    a.load_competitors(str(kniga_path))
    assert a.boot()["weight"] == "price"       # без продаж — веса по цене
    snap = a.set_weight("sales")
    assert "error" in snap and snap["weight"] == "price"


def test_filters_and_search(api):
    snap = api.set_filters({"selected": {"brand": ["Thomas"]}})
    assert snap["table"]["total"] == 3
    snap = api.set_filters({"search": "ROMA"})
    assert 3 <= snap["table"]["total"] <= 4
    snap = api.set_filters({"search": "20671, 925595"})   # список через запятую
    assert snap["table"]["total"] == 2
    snap = api.reset_filters()
    assert snap["table"]["total"] == 12
    assert snap["filters"]["activeCount"] == 0


def test_complete_only_filter(api):
    snap = api.set_filters({"complete_only": True})
    assert snap["table"]["total"] == 8
    skus = {r["sku"] for r in snap["table"]["rows"]}
    assert "20671" not in skus


def test_table_sort_and_paging(api):
    snap = api.set_table(sort_field="price")
    assert snap["table"]["sort"] == "price" and snap["table"]["asc"] is True
    first = snap["table"]["rows"][0]
    assert first["price"] == min(r["price"] for r in snap["table"]["rows"])
    snap = api.set_table(sort_field="price")   # повторный клик — обратный порядок
    assert snap["table"]["asc"] is False


def test_save_mapping_applies_and_persists(api, tmp_path):
    snap = api.save_mapping("Monobank", {"10": "5", "7": "4", "18": "9"})
    assert "error" not in snap
    assert snap["mapping"]["active"] and "Monobank" in snap["mapping"]["badge"]
    row = _row(snap, "20671")
    assert row["pay_comfy"] == 5 and row["dev_bank"] == -1
    assert (tmp_path / "mapping.json").exists()
    # другой банк маппинг не трогает
    privat = _row(api.set_bank("ПриватБанк"), "20671")
    assert privat["pay_comfy"] == 10


def test_save_mapping_validates(api):
    snap = api.save_mapping("Monobank", {"10": "abc"})
    assert "error" in snap


def test_save_apple_overrides_csv(api):
    snap = api.save_apple("Thomas", {"Monobank": "3", "ПриватБанк": "",
                                     "ПУМБ": ""})
    assert "error" not in snap
    assert snap["apple"]["active"]
    row = _row(api.set_bank("Monobank"), "20671")
    assert row["pay_comfy"] == 3 and row["dev_bank"] == -3
    # пустые поля → банк остаётся на значении из CSV
    privat = _row(api.set_bank("ПриватБанк"), "20671")
    assert privat["pay_comfy"] == 10


def test_export_to_writes_workbook(api, tmp_path):
    out = tmp_path / "report.xlsx"
    snap = api.export_to(str(out), all_banks=False)
    assert "error" not in snap and snap["saved"]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Сводка", "Срок и платёж", "Данные"]

    snap = api.export_to(str(tmp_path / "multi.xlsx"), all_banks=True)
    wb = openpyxl.load_workbook(tmp_path / "multi.xlsx")
    assert len(wb.sheetnames) == 3 * (len(api.dataset.banks) + 1)


def test_load_dropped_detects_file_kind(tmp_path, kniga_path, sales_path):
    a = Api(mapping_path=tmp_path / "m.json", override_path=tmp_path / "o.json")
    snap = a.load_dropped(str(kniga_path))     # csv конкурентов
    assert snap["loaded"] is True
    snap = a.load_dropped(str(sales_path))     # csv продаж → фолбэк
    assert snap["hasSales"] is True
    assert "error" not in snap


def test_avg_view_in_snapshot(api):
    snap = api.boot()
    view = snap["avgView"]
    assert view is not None
    assert view["retailers"][0]["name"] == "Comfy"
    assert view["retailers"][1]["name"] == "Все конкуренты"
    assert view["overall"]["category"] == "Вся выборка"
    rows = {r["category"]: r for r in view["rows"]}
    row = rows["Пилосос традиційний"]
    avg_price = (6649 + 9599 + 9599) / 3
    assert row["count"] == 3
    assert row["terms"]["pay_comfy"] == pytest.approx(10)
    assert row["payments"]["pay_comfy"] == pytest.approx(avg_price / 10)
    assert row["terms"]["pay::rozetka.com.ua"] == pytest.approx(3.5)
    assert row["payments"]["pay::rozetka.com.ua"] == pytest.approx(9599 / 3.5)
    assert row["terms"]["__market__"] == pytest.approx(5.0)       # все конкуренты
    assert row["payments"]["__market__"] == pytest.approx(9009 / 5)
    json.dumps(snap, allow_nan=False)

    # разрез следует за выбранным банком
    privat = api.set_bank("ПриватБанк")["avgView"]
    privat_row = {r["category"]: r for r in privat["rows"]}["Пилосос традиційний"]
    assert privat_row["terms"]["pay::foxtrot.com.ua"] == pytest.approx(15)


def test_kpi_values_match_metrics(api):
    snap = api.boot()
    comfy_card = next(c for c in snap["kpi"] if c["title"].startswith("Comfy"))
    assert comfy_card["accent"] is True
    value = float(comfy_card["value"])
    assert 0 < value < 30
