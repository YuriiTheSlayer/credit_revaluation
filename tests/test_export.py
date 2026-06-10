"""Тесты экспорта: структура листов, формулы итогов, сверка с эталоном."""

import openpyxl
import pandas as pd
import pytest

from core.model import ALL_BANKS, Dataset
from export.excel import SHEET_PRICE, SHEET_SALES, default_filename, export_report
from parsers.competitors import parse_competitors_csv
from parsers.sales import parse_sales
from test_metrics import load_reference_sales

ATTR_DEFAULTS = {
    "name": "", "brand": "", "subcategory": "", "business": "", "attr_rma": "",
    "supplier_item": "", "stock_comfy": "", "stock_sp": "", "discount": None,
    "margin": None, "step": None, "assort_type": "", "assort_type_global": "",
}
REF_NAMES = {
    "foxtrot.com.ua": "Foxtrot",
    "rozetka.com.ua": "Rozetka",
    "epicentrk.ua": "Epicentr",
}


def _reference_wide(example_xlsx_path) -> pd.DataFrame:
    """Широкая таблица из данных эталона + недостающие атрибуты-пустышки."""
    df, _ = load_reference_sales(example_xlsx_path)
    for col, value in ATTR_DEFAULTS.items():
        df[col] = value
    comp_cols = [f"pay::{d}" for d in REF_NAMES]
    df["comp_max_bank"] = df[comp_cols].max(axis=1)
    df["dev_bank"] = df["pay_comfy"] - df["comp_max_bank"]
    return df


@pytest.fixture
def dataset(kniga_path, sales_path) -> Dataset:
    return Dataset(parse_competitors_csv(kniga_path)).attach_sales(parse_sales(sales_path))


def _find_label_cell(ws, label):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == label:
                return cell
    return None


def test_default_filename():
    from datetime import date
    name = default_filename("Monobank", date(2026, 6, 10))
    assert name == "report_Monobank_2026-06-10.xlsx"


def test_reference_data_roundtrip(tmp_path, example_xlsx_path):
    """Экспорт данных эталона воспроизводит его итоги (±0.01) и структуру."""
    wide = _reference_wide(example_xlsx_path)
    out = tmp_path / "ref.xlsx"
    export_report(out, [("Monobank", wide)], REF_NAMES, has_sales=True)

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb.sheetnames == [SHEET_SALES, SHEET_PRICE]
    ws = wb[SHEET_SALES]

    # подписи ритейлеров под итогами
    comfy_label = _find_label_cell(ws, "Comfy")
    assert comfy_label is not None
    for disp in REF_NAMES.values():
        assert _find_label_cell(ws, disp) is not None

    # итог Comfy — в той же колонке строкой выше подписи; сверка с эталоном
    totals_row = comfy_label.row - 1
    comfy_total = ws.cell(row=totals_row, column=comfy_label.column).value
    assert comfy_total == pytest.approx(16.41, abs=0.01)
    fox_label = _find_label_cell(ws, "Foxtrot")
    assert ws.cell(row=totals_row, column=fox_label.column).value == pytest.approx(14.49, abs=0.01)
    roz_label = _find_label_cell(ws, "Rozetka")
    assert ws.cell(row=totals_row, column=roz_label.column).value == pytest.approx(15.46, abs=0.01)
    epi_label = _find_label_cell(ws, "Epicentr")
    assert ws.cell(row=totals_row, column=epi_label.column).value == pytest.approx(14.60, abs=0.01)

    # живые формулы: структура и SUMPRODUCT
    wb_f = openpyxl.load_workbook(out, data_only=False)
    ws_f = wb_f[SHEET_SALES]
    headers = {c.value: c.column for c in ws_f[2]}
    structure_col = headers["Структура"]
    first_structure = ws_f.cell(row=3, column=structure_col).value
    assert isinstance(first_structure, str) and first_structure.startswith("=")
    assert "/" in first_structure
    comfy_formula = ws_f.cell(row=totals_row, column=comfy_label.column).value
    assert isinstance(comfy_formula, str) and comfy_formula.startswith("=SUMPRODUCT(")
    total_sales_formula = ws_f.cell(row=totals_row, column=headers["Продажи"]).value
    assert total_sales_formula.startswith("=SUM(")


def test_export_structure_from_fixtures(tmp_path, dataset):
    """Сквозной сценарий: Книга31 + продажи → экспорт → проверка структуры."""
    bank = "Monobank"
    wide = dataset.wide(bank)
    out = tmp_path / "report.xlsx"
    export_report(
        out, [(bank, wide)], dataset.competitor_names(),
        has_sales=True, filters_desc="без фильтров",
    )

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [SHEET_SALES, SHEET_PRICE]
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        assert bank in str(ws["A1"].value)              # банк в шапке
        assert ws.freeze_panes == "C3"                  # заморозка шапки
        assert ws.auto_filter.ref is not None           # автофильтр
        headers = [c.value for c in ws[2]]
        assert "КодТовара" in headers and "Структура" in headers
        assert "Кол-во Платежей Foxtrot" in headers
        assert "Кол-во Платежей Rozetka" in headers
        assert len(list(ws.conditional_formatting)) > 0   # подсветка отклонений
    assert "Продажи" in [c.value for c in wb[SHEET_SALES][2]]
    assert "Продажи" not in [c.value for c in wb[SHEET_PRICE][2]]

    # у каждой категории своя итоговая строка с подписью Comfy
    ws = wb[SHEET_SALES]
    comfy_labels = [c for row in ws.iter_rows() for c in row if c.value == "Comfy"]
    n_categories = wide["category"].nunique()
    assert len(comfy_labels) == n_categories


def test_multibank_export_sheet_names(tmp_path, dataset):
    frames = [(b, dataset.wide(b)) for b in dataset.banks + [ALL_BANKS]]
    out = tmp_path / "multi.xlsx"
    export_report(out, frames, dataset.competitor_names(), has_sales=True)
    wb = openpyxl.load_workbook(out)
    assert len(wb.sheetnames) == 2 * len(frames)
    assert all(len(n) <= 31 for n in wb.sheetnames)
    assert any("Monobank" in n for n in wb.sheetnames)
    assert any("ПУМБ" in n for n in wb.sheetnames)


def test_export_without_sales_only_price_sheet(tmp_path, dataset):
    wide = dataset.wide("ПУМБ")
    out = tmp_path / "no_sales.xlsx"
    export_report(out, [("ПУМБ", wide)], dataset.competitor_names(), has_sales=False)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [SHEET_PRICE]


def test_export_empty_frame_writes_stub(tmp_path, dataset):
    wide = dataset.wide("Monobank").iloc[0:0]
    out = tmp_path / "empty.xlsx"
    export_report(out, [("Monobank", wide)], dataset.competitor_names(), has_sales=True)
    wb = openpyxl.load_workbook(out)
    ws = wb[SHEET_SALES]
    assert "Нет данных" in str(ws["A3"].value)
