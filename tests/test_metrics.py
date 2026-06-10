"""Сверка SUMPRODUCT с эталоном claude_example.xlsx + свойства перенормировки."""

import openpyxl
import pandas as pd
import pytest

from core import metrics
from core.model import ALL_BANKS, Dataset, Filters
from parsers.competitors import parse_competitors_csv, parse_number
from parsers.sales import parse_sales

COMP_COLS = ["pay::foxtrot.com.ua", "pay::rozetka.com.ua", "pay::epicentrk.ua"]
PAY_COLS = ["pay_comfy"] + COMP_COLS


def load_reference_sales(path):
    """Лист «Ср. срок на продажи» эталона → (df, ожидаемые итоги)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Ср. срок на продажи"]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=37, values_only=False):
        cells = {c.column_letter: c.value for c in r}
        rows.append({
            "sku": str(cells["A"]),
            "category": cells["E"],
            "sales": float(cells["L"]),
            "price": parse_number(cells["N"]),
            "pay_comfy": float(cells["P"]),
            "pay::foxtrot.com.ua": float(cells["V"]),
            "pay::rozetka.com.ua": float(cells["X"]),
            "pay::epicentrk.ua": float(cells["Y"]),
        })
    expected = {
        "pay_comfy": ws["P38"].value,
        "pay::foxtrot.com.ua": ws["V38"].value,
        "pay::rozetka.com.ua": ws["X38"].value,
        "pay::epicentrk.ua": ws["Y38"].value,
    }
    return pd.DataFrame(rows), expected


def load_reference_price(path):
    """Лист «Средний срок на стоимость» эталона → (df, ожидаемые итоги)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Средний срок на стоимость"]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=37):
        cells = {c.column_letter: c.value for c in r}
        rows.append({
            "sku": str(cells["A"]),
            "category": cells["E"],
            "price": float(cells["L"]),
            "pay_comfy": float(cells["O"]),
            "pay::foxtrot.com.ua": float(cells["U"]),
            "pay::rozetka.com.ua": float(cells["W"]),
            "pay::epicentrk.ua": float(cells["X"]),
        })
    expected = {
        "pay_comfy": ws["O38"].value,
        "pay::foxtrot.com.ua": ws["U38"].value,
        "pay::rozetka.com.ua": ws["W38"].value,
        "pay::epicentrk.ua": ws["X38"].value,
    }
    return pd.DataFrame(rows), expected


def test_sumproduct_matches_reference_sales_sheet(example_xlsx_path):
    """Критерий приёмки: совпадение с эталоном (16.41/14.49/15.46/14.60) до 0.01."""
    df, expected = load_reference_sales(example_xlsx_path)
    terms = metrics.weighted_terms(df, "sales", PAY_COLS)
    row = terms.loc["Смартфон"]
    assert row["pay_comfy"] == pytest.approx(expected["pay_comfy"], abs=0.01)
    assert row["pay_comfy"] == pytest.approx(16.41, abs=0.01)
    assert row["pay::foxtrot.com.ua"] == pytest.approx(14.49, abs=0.01)
    assert row["pay::rozetka.com.ua"] == pytest.approx(15.46, abs=0.01)
    assert row["pay::epicentrk.ua"] == pytest.approx(14.60, abs=0.01)
    for col in PAY_COLS:
        assert row[col] == pytest.approx(expected[col], abs=0.01)


def test_sumproduct_matches_reference_price_sheet(example_xlsx_path):
    df, expected = load_reference_price(example_xlsx_path)
    terms = metrics.weighted_terms(df, "price", PAY_COLS)
    row = terms.loc["Смартфон"]
    for col in PAY_COLS:
        assert row[col] == pytest.approx(expected[col], abs=0.01)


def test_structure_sums_to_one_per_category():
    df = pd.DataFrame({
        "category": ["A", "A", "B", "B", "B"],
        "sales": [10, 30, 5, 0, 15],
    })
    share = metrics.structure(df, "sales")
    sums = share.groupby(df["category"]).sum()
    assert sums["A"] == pytest.approx(1)
    assert sums["B"] == pytest.approx(1)
    assert share.iloc[0] == pytest.approx(0.25)


def test_missing_payments_excluded_with_renormalization():
    """SKU без данных у ритейлера не занижает его среднюю (веса перенормируются)."""
    df = pd.DataFrame({
        "category": ["A", "A"],
        "sales": [75.0, 25.0],
        "pay_x": [10.0, None],
    })
    terms = metrics.weighted_terms(df, "sales", ["pay_x"])
    assert terms.loc["A", "pay_x"] == pytest.approx(10.0)   # не 7.5


def test_zero_weight_base_gives_nan():
    df = pd.DataFrame({"category": ["A"], "sales": [0.0], "pay_x": [10.0]})
    terms = metrics.weighted_terms(df, "sales", ["pay_x"])
    assert pd.isna(terms.loc["A", "pay_x"])


def test_filters_renormalize_weights(kniga_path, sales_path):
    """При активных фильтрах веса пересчитываются по отфильтрованному набору."""
    dataset = Dataset(parse_competitors_csv(kniga_path)).attach_sales(parse_sales(sales_path))
    wide = dataset.wide("Monobank")
    flt = Filters(brand={"Thomas"})
    filtered = flt.apply(wide)
    assert 0 < len(filtered) < len(wide)
    share = metrics.structure(filtered, "sales")
    sums = share.groupby(filtered["category"]).sum()
    for cat, s in sums.items():
        base = filtered[filtered["category"] == cat]["sales"].fillna(0).sum()
        if base > 0:
            assert s == pytest.approx(1), f"веса категории {cat} должны давать 1"


def test_overall_terms_and_win_share(kniga_path, sales_path):
    dataset = Dataset(parse_competitors_csv(kniga_path)).attach_sales(parse_sales(sales_path))
    wide = dataset.wide(ALL_BANKS)
    pay_cols = ["pay_comfy"] + [f"pay::{c}" for c in dataset.competitors]
    overall = metrics.overall_terms(wide, "sales", pay_cols)
    assert overall["pay_comfy"] > 0
    ws = metrics.win_share(wide)
    assert ((ws >= 0) & (ws <= 1)).all()
